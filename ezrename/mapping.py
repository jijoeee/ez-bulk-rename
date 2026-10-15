"""Load a rename list (old name -> new name) from Excel or CSV.

The list is the first two filled-in columns of the first sheet: column 1 is the
current file name (with or without its extension), column 2 the new name. A
header row such as "Old name | New name" or "Staff ID | Name" is skipped.

No UI imports. Problems are RenameError sentences.
"""

from __future__ import annotations

import csv
import os
import re

from .rules import RenameError

# A first row is a header if its cells look like column titles, not file names.
_HEADER_WORDS = ("old", "new", "current", "from", "to", "name", "nama", "file", "id", "staff", "lama", "baru")


def _looks_like_header(a: str, b: str) -> bool:
    """"Old name | New name" or "Staff ID | Nama" is a header; "S001 | John Tan" is data.

    File names almost always have a digit or a dot; column titles don't.
    """
    words = set(re.findall(r"[a-z]+", (a + " " + b).lower()))
    return bool(words & set(_HEADER_WORDS)) and not re.search(r"[\d.]", a)


def _rows_from_csv(path: str) -> list[list[str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(f, dialect)]


def _rows_from_excel(path: str) -> list[list[str]]:
    from openpyxl import load_workbook  # imported here: only needed when a workbook is loaded

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise RenameError(f"{os.path.basename(path)} can't be opened as an Excel workbook. Save it as "
                          ".xlsx (or .csv) and try again.") from None
    try:
        sheet = wb.worksheets[0]
        rows = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        wb.close()
    return rows


def load_mapping(path: str) -> tuple[dict[str, str], int]:
    """Return ({lower-case old name: new name}, number of duplicate old names ignored)."""
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            rows = _rows_from_excel(path)
        elif ext in (".csv", ".txt"):
            rows = _rows_from_csv(path)
        else:
            raise RenameError(f"{name}: use an Excel workbook (.xlsx) or a CSV file.")
    except RenameError:
        raise
    except UnicodeDecodeError:
        raise RenameError(f"{name} is not a plain-text CSV. Save it from Excel as \"CSV UTF-8\".") from None
    except OSError as exc:
        raise RenameError(f"Can't open {name}: {exc.strerror or 'disk error'}. Close it in Excel and "
                          "try again.") from None

    pairs = []
    for row in rows:
        cells = [c.strip() for c in row if c is not None and str(c).strip()]
        if len(cells) >= 2:
            pairs.append((cells[0], cells[1]))
    if pairs and _looks_like_header(*pairs[0]):
        pairs = pairs[1:]
    if not pairs:
        raise RenameError(f"{name} has no rows with two filled-in columns (current name, new name).")

    mapping: dict[str, str] = {}
    duplicates = 0
    for old, new in pairs:
        key = old.lower()
        if key in mapping:
            duplicates += 1
            continue
        mapping[key] = new
    return mapping, duplicates
